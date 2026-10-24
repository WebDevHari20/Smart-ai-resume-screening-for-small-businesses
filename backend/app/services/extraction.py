"""Stage 1: resume upload -> raw text, hyperlinks, OCR fallback, hidden-text detection."""
import io
import re

URL_RE = re.compile(r"https?://[^\s)>\]]+", re.IGNORECASE)
ALLOWED_EXT = {".pdf", ".docx", ".txt"}


class ResumeError(Exception):
    pass


def _is_white(color) -> bool:
    if color is None:
        return False
    if isinstance(color, (int, float)):
        return color >= 0.98
    try:
        vals = [float(v) for v in color]
    except (TypeError, ValueError):
        return False
    return bool(vals) and all(v >= 0.98 for v in vals) and len(vals) in (1, 3)


def _pdf(content: bytes):
    import pdfplumber

    pages, links, hidden, ocr = [], [], 0, False
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
            try:
                links += [h["uri"] for h in page.hyperlinks if h.get("uri")]
            except Exception:
                pass
            for ch in page.chars:
                if ch["text"].strip() and (_is_white(ch.get("non_stroking_color")) or ch.get("size", 10) < 2):
                    hidden += 1
    text = "\n".join(pages).strip()
    if len(text) < 50:  # scanned document -> OCR fallback (optional deps)
        try:
            import pytesseract
            from pdf2image import convert_from_bytes

            text = "\n".join(pytesseract.image_to_string(img) for img in convert_from_bytes(content, dpi=200)).strip()
            ocr = True
        except Exception:
            pass
    return text, links, hidden, ocr


def _docx(content: bytes):
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    links = []
    for rel in doc.part.rels.values():
        if "hyperlink" in rel.reltype:
            ref = getattr(rel, "target_ref", None)
            if ref:
                links.append(ref)
    hidden = sum(1 for p in doc.paragraphs for r in p.runs if r.font.hidden and r.text.strip())
    return "\n".join(parts).strip(), links, hidden, False


def extract_text(filename: str, content: bytes) -> dict:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXT:
        raise ResumeError(f"Unsupported file type '{ext}'. Allowed: PDF, DOCX, TXT")
    try:
        if ext == ".pdf":
            text, links, hidden, ocr = _pdf(content)
        elif ext == ".docx":
            text, links, hidden, ocr = _docx(content)
        else:
            text, links, hidden, ocr = content.decode("utf-8", errors="ignore").strip(), [], 0, False
    except Exception as e:  # corrupt / encrypted files
        raise ResumeError(f"Could not read file: {e}")
    if len(text) < 30:
        raise ResumeError("No readable text found (empty or unreadable document)")
    links = list(dict.fromkeys(links + URL_RE.findall(text)))
    return {"text": text, "links": links, "hidden_chars": hidden, "used_ocr": ocr}
